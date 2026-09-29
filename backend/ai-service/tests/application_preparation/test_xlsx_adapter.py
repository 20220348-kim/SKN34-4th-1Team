import asyncio
import base64
from copy import copy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side, Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table
from openpyxl.formatting.rule import CellIsRule

from app.application_preparation.xlsx_adapter import XlsxDocumentAdapter, cell_value
from app.application_preparation.document_contract import (
    DocumentError, EditOperation, GenerateDocumentRequest, MapDocumentRequest,
    MappingSelection, PlanSelection, ENGINES, PIPELINE_VERSION, MAP_VERSION,
    digest, validate_plan, validate_mapping,
)
from app.application_preparation.agent import ApplicationPreparationAgent


def fixture(*, protection=False, validation_source='"서울,부산"', extra=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '신청 정보'
    sheet['A1'] = '지원 신청서'
    sheet.merge_cells('A1:D1')
    labels = {2: '기업명', 3: '사업자등록번호', 4: '합계', 5: '지원금액', 6: '이메일',
              7: '주소', 8: '숨김 행', 9: '숨김 열', 10: '지역', 11: '신청일', 12: '미해결 선택'}
    for row, label in labels.items():
        sheet.cell(row, 1, label)
        cell = sheet.cell(row, 2)
        cell.border = Border(bottom=Side(style='thin', color='000000'))
        cell.font = Font(name='Arial', bold=False, color='FF0000')
        cell.fill = PatternFill('solid', fgColor='FFFFAA')
        cell.alignment = Alignment(wrap_text=True)
    sheet['B3'].number_format = '0000000000'
    sheet['B4'] = '=SUM(B5:B6)'
    sheet['B5'].number_format = '#,##0'
    sheet['B11'].number_format = 'yyyy-mm-dd'
    sheet['B13'] = '이미 존재하는 값'
    sheet['C13'] = 123
    sheet['E13'] = True
    sheet['F13'] = datetime(2026, 9, 27)
    sheet.merge_cells('B7:C7')
    sheet.row_dimensions[8].hidden = True
    sheet.column_dimensions['D'].hidden = True
    sheet['D9'].border = copy(sheet['B2'].border)
    sheet.freeze_panes = 'B2'
    sheet.row_dimensions[2].height = 24
    sheet.column_dimensions['B'].width = 28
    sheet['J20'].border = copy(sheet['B2'].border)
    for address, source in [('B10', validation_source), ('B12', 'INDIRECT("missing")')]:
        rule = DataValidation(type='list', formula1=source)
        sheet.add_data_validation(rule)
        rule.add(address)
    sheet.conditional_formatting.add('B5', CellIsRule(operator='greaterThan', formula=['0'],
        fill=PatternFill('solid', fgColor='00FF00')))
    workbook.create_sheet('목록').sheet_state = 'veryHidden'
    workbook['목록']['A1'], workbook['목록']['A2'] = '서울', '부산'
    if protection:
        sheet.protection.sheet = True
        sheet['B3'].protection = Protection(locked=False)
    if extra:
        extra(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def source(tmp_path, data=None):
    path = tmp_path / 'source.xlsx'
    path.write_bytes(data or fixture())
    return path


def inspect(path):
    return asyncio.run(XlsxDocumentAdapter().inspect(path))


def target(document, address, sheet='신청 정보'):
    return next(t for t in document.targets if t.nativeLocator['cellAddress'] == address
                and t.nativeLocator['sheetName'] == sheet)


def request(data, value='테스트기업'):
    return GenerateDocumentRequest(sourceBase64=base64.b64encode(data).decode(), sourceSha256=digest(data),
        format='xlsx', answerRevision=1, scope='지원 신청서',
        facts=[{'id': 'company:name', 'label': '기업명', 'value': value}])


def operation(t, *, kind='input'):
    return EditOperation(targetId=t.targetId, operation=kind, expectedText=t.currentText,
        start=0, end=len(t.currentText), valueRef='company:name', box=None, reason='Verified blank input')


def plan(path, document, address='B2', value='테스트기업', kind='input'):
    t = target(document, address)
    return validate_plan(request(path.read_bytes(), value), document, PlanSelection(
        operations=[operation(t, kind=kind)], unresolvedTargets=[], scopeTargetIds=[t.targetId]))


def test_native_cell_metadata_and_sheet_structure(tmp_path):
    path = source(tmp_path)
    document = inspect(path)
    t = target(document, 'B2')
    assert t.targetId == 'xlsx:s:%EC%8B%A0%EC%B2%AD%20%EC%A0%95%EB%B3%B4:c:B2'
    assert t.editable and t.nativeLocator['fieldLabels'] == ['기업명']
    assert t.nativeLocator['row'] == 2 and t.nativeLocator['column'] == 2
    assert t.nativeLocator['value'] is None and t.nativeLocator['displayValue'] == ''
    info = document.workbookMetadata['sheets'][0]
    assert info['freezePanes'] == 'B2' and info['formulaCount'] == 1
    assert info['mergedRanges'] == ['A1:D1', 'B7:C7']
    assert len(info['dataValidations']) == 2
    assert document.workbookMetadata['sheets'][1]['state'] == 'veryHidden'
    assert target(document, 'C13').nativeLocator['value'] == 123
    assert target(document, 'E13').nativeLocator['dataType'] == 'b'
    assert target(document, 'F13').nativeLocator['dataType'] == 'd'


@pytest.mark.parametrize('address,reason', [('B4', 'FORMULA_CELL'), ('C7', 'MERGED_CHILD'),
    ('B8', 'HIDDEN_CELL'), ('D9', 'HIDDEN_CELL'), ('J20', 'AMBIGUOUS_BLANK_CELL'),
    ('B12', 'UNRESOLVED_OPTION'), ('B13', 'NONEMPTY_CELL')])
def test_unsafe_targets_are_readonly_and_validator_rejects_them(tmp_path, address, reason):
    path = source(tmp_path)
    document = inspect(path)
    t = target(document, address)
    assert not t.editable and t.unsupportedReason == reason
    with pytest.raises(DocumentError):
        plan(path, document, address, kind='set_field')


def test_merged_master_is_the_only_address_that_can_be_written(tmp_path):
    path = source(tmp_path)
    document = inspect(path)
    master = target(document, 'B7')
    assert master.editable and master.nativeLocator['mergedMaster'] == 'B7'
    output, result = asyncio.run(XlsxDocumentAdapter().apply(path, document, plan(path, document, 'B7'),
        {'company:name': '테스트기업'}))
    workbook = load_workbook(BytesIO(output))
    assert workbook.active['B7'].value == '테스트기업' and workbook.active['C7'].value is None
    assert 'B7:C7' in list(map(str, workbook.active.merged_cells.ranges))
    assert result['formulas'] == result['dataValidation'] == 'PASSED'
    workbook.close()


def test_protected_sheet_only_unlocked_inputs_can_be_written(tmp_path):
    document = inspect(source(tmp_path, fixture(protection=True)))
    assert target(document, 'B2').unsupportedReason == 'PROTECTED_CELL'
    assert target(document, 'B3').editable
    assert not target(document, 'A1', '목록').editable


def test_workbook_structure_protection_is_not_bypassed(tmp_path):
    document = inspect(source(tmp_path, fixture(extra=lambda w: setattr(w.security, 'lockStructure', True))))
    assert target(document, 'B2').unsupportedReason == 'PROTECTED_CELL'


@pytest.mark.parametrize('value', ['테스트기업', '=SUM(1,2)', '001-12-34567', 'line1\nline2'])
def test_write_preserves_all_unedited_parts_formulas_styles_and_properties(tmp_path, value):
    path = source(tmp_path)
    original = path.read_bytes()
    document = inspect(path)
    output, result = asyncio.run(XlsxDocumentAdapter().apply(path, document, plan(path, document, value=value),
        {'company:name': value}))
    assert path.read_bytes() == original
    workbook = load_workbook(BytesIO(output))
    assert workbook.active['B2'].value == value and workbook.active['B2'].data_type == 's'
    assert workbook.active['B4'].value == '=SUM(B5:B6)'
    assert workbook.active['B13'].value == '이미 존재하는 값'
    assert workbook.active['B2'].font.color.rgb == '00FF0000'
    assert workbook.active.row_dimensions[2].height == 24
    assert workbook.active.column_dimensions['B'].width == 28
    assert workbook['목록'].sheet_state == 'veryHidden'
    workbook.close()
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(output)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != 'xl/worksheets/sheet1.xml':
                assert before.read(name) == after.read(name)
    assert result['reopened'] is True and result['verified'] == 1
    assert result['unchangedParts'] == 'PASSED'


@pytest.mark.parametrize('source_list', ['"서울,부산"', "'목록'!$A$1:$A$2"])
def test_dropdown_accepts_only_resolved_native_options(tmp_path, source_list):
    path = source(tmp_path, fixture(validation_source=source_list))
    document = inspect(path)
    assert target(document, 'B10').nativeLocator['dataValidation']['allowedValues'] == ['서울', '부산']
    with pytest.raises(DocumentError) as error:
        plan(path, document, 'B10', value='Seoul')
    assert error.value.reason == 'UNRESOLVED_OPTION'
    output, _ = asyncio.run(XlsxDocumentAdapter().apply(path, document,
        plan(path, document, 'B10', value='서울'), {'company:name': '서울'}))
    workbook = load_workbook(BytesIO(output))
    assert workbook.active['B10'].value == '서울'
    assert len(workbook.active.data_validations.dataValidation) == 2
    workbook.close()


@pytest.mark.parametrize('address,value,expected', [('B3', '0011234567', '0011234567'),
    ('B5', '1000000', 1000000), ('B11', '2026-09-27', datetime(2026, 9, 27))])
def test_value_types_follow_format_and_identifiers_keep_zeroes(tmp_path, address, value, expected):
    path = source(tmp_path)
    document = inspect(path)
    output, _ = asyncio.run(XlsxDocumentAdapter().apply(path, document,
        plan(path, document, address, value=value), {'company:name': value}))
    workbook = load_workbook(BytesIO(output))
    assert workbook.active[address].value == expected
    assert workbook.active[address].number_format == target(document, address).nativeLocator['numberFormat']
    workbook.close()


@pytest.mark.parametrize('address,value', [('B5', '1,000,000원'), ('B5', 'NaN'), ('B5', '1e999'),
    ('B11', '내일'), ('B11', '2026-09-27T00:00:00+09:00')])
def test_ambiguous_numeric_and_date_values_fail_explicitly(tmp_path, address, value):
    path = source(tmp_path)
    with pytest.raises(DocumentError):
        plan(path, inspect(path), address, value=value)


def test_source_hash_and_direct_adapter_guard_reject_stale_or_forged_plans(tmp_path):
    path = source(tmp_path)
    document = inspect(path)
    valid = plan(path, document)
    forged = valid.model_copy(deep=True)
    forged.operations[0].targetId = target(document, 'B4').targetId
    forged.operations[0].expectedText = '=SUM(B5:B6)'
    forged.scopeTargetIds = [forged.operations[0].targetId]
    with pytest.raises(DocumentError):
        asyncio.run(XlsxDocumentAdapter().apply(path, document, forged, {'company:name': 'テスト'}))
    path.write_bytes(fixture(protection=True))
    with pytest.raises(DocumentError) as error:
        asyncio.run(XlsxDocumentAdapter().apply(path, document, valid, {'company:name': 'test'}))
    assert error.value.code == 'APPLICATION_DOCUMENT_SOURCE_CHANGED'


def mutated(data, name, content):
    buffer = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(buffer, 'w') as output:
        for entry in original.infolist():
            output.writestr(entry, content if entry.filename == name else original.read(entry))
        if name not in original.namelist():
            output.writestr(name, content)
    return buffer.getvalue()


@pytest.mark.parametrize('part', ['xl/vbaProject.bin', 'xl/externalLinks/link1.xml',
    'xl/embeddings/object1.bin', 'xl/drawings/drawing1.xml', '_xmlsignatures/signature.xml'])
def test_macros_links_and_unsupported_objects_are_rejected(tmp_path, part):
    data = mutated(fixture(), part, b'<unsupported/>')
    with pytest.raises(DocumentError):
        inspect(source(tmp_path, data))


@pytest.mark.parametrize('payload', [b'<!DOCTYPE a [<!ENTITY x "bad">]><a/>',
    '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE a [<!ENTITY x "bad">]><a/>'.encode('utf-16')])
def test_xml_entities_are_rejected_in_any_part(tmp_path, payload):
    with pytest.raises(DocumentError) as error:
        inspect(source(tmp_path, mutated(fixture(), 'docProps/core.xml', payload)))
    assert error.value.reason == 'XLSX_UNSAFE_XML'


def test_zip_entry_count_and_expansion_budget(tmp_path, monkeypatch):
    import app.application_preparation.xlsx_adapter as module
    data = fixture()
    monkeypatch.setattr(module, 'MAX_BYTES', len(data) + 10)
    with pytest.raises(DocumentError) as error:
        inspect(source(tmp_path, data))
    assert error.value.code == 'APPLICATION_DOCUMENT_LIMIT_EXCEEDED'


def test_mapping_context_contains_only_editable_candidates_and_excludes_ground_truth(tmp_path):
    path = source(tmp_path)
    document = inspect(path)
    mapping = MapDocumentRequest(sourceBase64=base64.b64encode(path.read_bytes()).decode(),
        sourceSha256=document.sourceSha256, format='xlsx', scope='지원 신청서',
        fields=[{'id': 'company:name', 'label': '기업명', 'guidance': '', 'required': True}])
    t = target(document, 'B2')
    selection = MappingSelection(bindings=[{'factId': 'company:name', 'targetId': t.targetId, 'box': None}],
        scopeTargetIds=[t.targetId], unmappedFieldIds=[])
    agent = ApplicationPreparationAgent(model=None, run_timeout_seconds=3)
    agent._invoke = AsyncMock(return_value=selection)
    asyncio.run(agent.map_document(mapping, document))
    payload = json.loads(agent._invoke.call_args.args[2][0]['text'])
    assert all(item['editable'] for item in payload['documentMap']['targets'])
    assert t.targetId in payload['fieldCandidates']['company:name']
    assert 'expectedTargetId' not in json.dumps(payload)
    validate_mapping(mapping, document, selection)
    unsafe = target(document, 'B4')
    bad = selection.model_copy(deep=True)
    bad.bindings[0].targetId = unsafe.targetId
    bad.scopeTargetIds = [unsafe.targetId]
    with pytest.raises(DocumentError):
        validate_mapping(mapping, document, bad)


def test_xlsx_engine_does_not_invalidate_existing_four_format_pipeline():
    import hashlib
    from app.application_preparation.document_contract import CONTRACT, PLAN_VERSION, KORDOC_VERSION
    previous = hashlib.sha256(json.dumps([CONTRACT, MAP_VERSION, PLAN_VERSION,
        {k: v for k, v in ENGINES.items() if k not in {'docx', 'xlsx'}}, KORDOC_VERSION],
        sort_keys=True).encode()).hexdigest()
    assert PIPELINE_VERSION == previous
    assert ENGINES['docx'] == 'govbiz/ooxml-native@2'
    assert ENGINES['xlsx'] == 'govbiz/xlsx-native@2+openpyxl-3.1.5'


def test_real_ai_router_map_generate_roundtrip_without_paid_model(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.application_preparation.router import get_service, router

    class Agent:
        async def map_document(self, req, document):
            t = target(document, 'B2')
            return MappingSelection(bindings=[{'factId': 'company:name', 'targetId': t.targetId, 'box': None}],
                scopeTargetIds=[t.targetId], unmappedFieldIds=[])

        async def plan_document(self, req, document):
            # 저장된 바인딩이 있으면 작성 계획은 결정적으로 만들어지므로 모델 계획은 호출되지 않아야 합니다.
            raise AssertionError('xlsx plan must be derived from saved bindings')

    monkeypatch.setenv('DOCUMENT_INTERNAL_TOKEN', 't' * 32)
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_service] = lambda: SimpleNamespace(agent=Agent())
    data = fixture()
    generation = request(data).model_dump()
    mapping = {**generation, 'facts': [], 'fields': [{'id': 'company:name', 'label': '기업명',
               'guidance': '', 'required': True}]}
    headers = {'Authorization': 'Bearer ' + 't' * 32}
    with TestClient(app) as client:
        configuration = client.get('/internal/v1/application-preparations/document/configuration').json()
        assert configuration['engineVersions']['xlsx'] == ENGINES['xlsx']
        mapped = client.post('/internal/v1/application-preparations/document/map', json=mapping, headers=headers)
        assert mapped.status_code == 200, mapped.text
        generation.update(bindings=mapped.json()['bindings'], scopeTargetIds=mapped.json()['scopeTargetIds'])
        response = client.post('/internal/v1/application-preparations/document/generate', json=generation, headers=headers)
        assert response.status_code == 200, response.text
    result = response.json()
    output = base64.b64decode(result['outputBase64'])
    assert result['outputSha256'] == digest(output) and result['verification']['reopened']
    workbook = load_workbook(BytesIO(output))
    assert workbook.active['B2'].value == '테스트기업'
    workbook.close()




def test_rich_text_heading_is_label_evidence_without_losing_its_runs(tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    def rich(workbook):
        workbook.active['A2'] = CellRichText([TextBlock(InlineFont(b=True), '기업'), '명'])
    path = source(tmp_path, fixture(extra=rich))
    document = inspect(path)
    assert target(document, 'B2').editable
    assert target(document, 'B2').nativeLocator['fieldLabels'] == ['기업명']
    output, _ = asyncio.run(XlsxDocumentAdapter().apply(path, document, plan(path, document),
        {'company:name': '테스트기업'}))
    workbook = load_workbook(BytesIO(output), rich_text=True)
    assert isinstance(workbook.active['A2'].value, CellRichText)
    assert str(workbook.active['A2'].value) == '기업명'
    workbook.close()


def test_xlsx_labels_do_not_bind_representative_to_team_member(tmp_path):
    from app.application_preparation.document_contract import mapping_label_matches
    document = inspect(source(tmp_path))
    t = target(document, 'B2').model_copy(deep=True)
    t.nativeLocator['fieldLabels'] = ['성명']
    assert not mapping_label_matches('신청자 성명(대표자)', t)
    t.nativeLocator['fieldLabels'] = ['신청자 성명(대표자)']
    assert mapping_label_matches('기업 / 신청자 성명(대표자)', t)



@pytest.mark.parametrize('number_format,value,expected', [('₩#,##0', '1000000', 1000000), ('0.0%', '12.5%', 0.125)])
def test_currency_and_percentage_keep_native_number_formats(tmp_path, number_format, value, expected):
    def formatted(workbook):
        workbook.active['B5'].number_format = number_format
    path = source(tmp_path, fixture(extra=formatted))
    document = inspect(path)
    output, _ = asyncio.run(XlsxDocumentAdapter().apply(path, document,
        plan(path, document, 'B5', value=value), {'company:name': value}))
    workbook = load_workbook(BytesIO(output))
    assert workbook.active['B5'].value == expected
    assert workbook.active['B5'].number_format == number_format
    workbook.close()


def test_defined_name_dropdown_is_resolved_from_its_native_source(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    def named(workbook):
        workbook.defined_names.add(DefinedName('Regions', attr_text="'목록'!$A$1:$A$2"))
    path = source(tmp_path, fixture(validation_source='Regions', extra=named))
    assert target(inspect(path), 'B10').nativeLocator['dataValidation']['allowedValues'] == ['서울', '부산']


def test_named_data_table_cells_are_not_form_inputs(tmp_path):
    def table(workbook):
        sheet = workbook.create_sheet('데이터')
        sheet.append(['항목', '값'])
        sheet.append(['기업명', None])
        sheet.append(['주소', None])
        sheet['B2'].border = Border(bottom=Side(style='thin'))
        sheet['B3'].border = copy(sheet['B2'].border)
        sheet.add_table(Table(displayName='SourceData', ref='A1:B3'))
    document = inspect(source(tmp_path, fixture(extra=table)))
    cell = target(document, 'B2', '데이터')
    assert not cell.editable and cell.unsupportedReason == 'DATA_TABLE'
    assert cell.nativeLocator['tableRegion'] == 'A1:B3'


def test_unknown_macro_part_name_is_rejected_by_content_type(tmp_path):
    data = fixture()
    with ZipFile(BytesIO(data)) as archive:
        types = archive.read('[Content_Types].xml').replace(b'</Types>',
            b'<Override PartName="/xl/custom.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
    data = mutated(mutated(data, '[Content_Types].xml', types), 'xl/custom.bin', b'not executed')
    with pytest.raises(DocumentError) as error:
        inspect(source(tmp_path, data))
    assert error.value.reason == 'XLSX_UNSUPPORTED_OBJECT_OR_EXTERNAL_LINK'


def test_comment_vml_and_extended_list_validation_are_preserved_without_losing_allowed_values(tmp_path):
    def note(workbook):
        workbook.active['B2'].comment = Comment('원문 주석', '작성자')
    data = fixture(extra=note)
    extension = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}">'
        '<x14:dataValidations xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
        'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" count="1">'
        '<x14:dataValidation type="list"><x14:formula1><xm:f>\'목록\'!$A$1:$A$2</xm:f></x14:formula1>'
        '<xm:sqref>B2</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>').encode()
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').replace(b'</worksheet>', extension + b'</worksheet>')
    path = source(tmp_path, mutated(data, 'xl/worksheets/sheet1.xml', xml))
    document = inspect(path)
    assert target(document, 'B2').nativeLocator['dataValidation']['allowedValues'] == ['서울', '부산']
    with pytest.raises(DocumentError):
        cell_value(target(document, 'B2'), '대전')
    output, _ = asyncio.run(XlsxDocumentAdapter().apply(path, document,
        plan(path, document, value='서울'), {'company:name': '서울'}))
    with ZipFile(BytesIO(path.read_bytes())) as before, ZipFile(BytesIO(output)) as after:
        for name in before.namelist():
            if name != 'xl/worksheets/sheet1.xml':
                assert before.read(name) == after.read(name)
        assert extension in after.read('xl/worksheets/sheet1.xml')


def test_vml_with_non_comment_control_remains_unsupported(tmp_path):
    data = fixture(extra=lambda workbook: setattr(workbook.active['B2'], 'comment', Comment('참고', '작성자')))
    with ZipFile(BytesIO(data)) as archive:
        vml_name = next(name for name in archive.namelist() if name.endswith('.vml'))
        vml = archive.read(vml_name).replace(b'ObjectType="Note"', b'ObjectType="Button"')
    with pytest.raises(DocumentError) as error:
        inspect(source(tmp_path, mutated(data, vml_name, vml)))
    assert error.value.reason == 'XLSX_UNSUPPORTED_OBJECT_OR_EXTERNAL_LINK'


def test_unknown_xlsx_extension_is_not_silently_discarded(tmp_path):
    data = fixture()
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').replace(b'</worksheet>',
            b'<extLst><ext uri="unknown"><unknown/></ext></extLst></worksheet>')
    with pytest.raises(DocumentError) as error:
        inspect(source(tmp_path, mutated(data, 'xl/worksheets/sheet1.xml', xml)))
    assert error.value.reason == 'XLSX_EXTENDED_VALIDATION_UNSUPPORTED'


def test_large_styled_blank_grid_keeps_labeled_inputs_without_target_overflow(tmp_path):
    def formatted(workbook):
        for row in range(21, 501):
            for column in range(1, 12):
                workbook.active.cell(row, column).border = Border(bottom=Side(style='thin'))
    document = inspect(source(tmp_path, fixture(extra=formatted)))
    assert len(document.targets) < 3000
    assert target(document, 'B2').editable


def test_duplicate_zip_entries_and_cell_budget_are_rejected(tmp_path):
    data = BytesIO(fixture())
    with pytest.warns(UserWarning), ZipFile(data, 'a') as archive:
        archive.writestr('xl/workbook.xml', archive.read('xl/workbook.xml'))
    with pytest.raises(DocumentError):
        inspect(source(tmp_path, data.getvalue()))
    data = fixture()
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').replace(b'ref="A1:J20"', b'ref="A1:XFD1048576"')
    with pytest.raises(DocumentError) as error:
        inspect(source(tmp_path, mutated(data, 'xl/worksheets/sheet1.xml', xml)))
    assert error.value.code == 'APPLICATION_DOCUMENT_LIMIT_EXCEEDED'


def test_grouped_hidden_columns_do_not_produce_candidates(tmp_path):
    path = source(tmp_path, fixture(extra=lambda w: w.active.column_dimensions.group('B', 'D', hidden=True)))
    assert target(inspect(path), 'B2').unsupportedReason == 'HIDDEN_CELL'



def test_hidden_label_does_not_make_a_visible_blank_cell_editable(tmp_path):
    path = source(tmp_path, fixture(extra=lambda w: setattr(w.active.column_dimensions['A'], 'hidden', True)))
    t = target(inspect(path), 'B2')
    assert not t.editable and t.unsupportedReason == 'AMBIGUOUS_BLANK_CELL'
    assert t.nativeLocator['fieldLabels'] == []


def test_signature_part_is_rejected_even_with_a_nonstandard_name(tmp_path):
    data = fixture()
    with ZipFile(BytesIO(data)) as archive:
        types = archive.read('[Content_Types].xml').replace(b'</Types>',
            b'<Override PartName="/custom/signature.xml" ContentType="application/vnd.openxmlformats-package.digital-signature-xmlsignature+xml"/></Types>')
    data = mutated(mutated(data, '[Content_Types].xml', types), 'custom/signature.xml', b'<Signature/>')
    with pytest.raises(DocumentError):
        inspect(source(tmp_path, data))



@pytest.mark.parametrize('address,value,expected', [('B2', '테스트기업', '테스트기업'),
    ('B5', '1000000', 1000000), ('B11', '2026-09-27', datetime(2026, 9, 27))])
def test_prompt_only_validation_is_preserved_without_inventing_option_restrictions(tmp_path, address, value, expected):
    def prompted(workbook):
        rule = DataValidation(prompt='원문 입력 안내', showInputMessage=True, allow_blank=True)
        workbook.active.add_data_validation(rule)
        rule.add(address)
    path = source(tmp_path, fixture(extra=prompted))
    document = inspect(path)
    t = target(document, address)
    assert t.editable and t.nativeLocator['dataValidation']['promptOnly']
    output, _ = asyncio.run(XlsxDocumentAdapter().apply(path, document,
        plan(path, document, address, value=value), {'company:name': value}))
    workbook = load_workbook(BytesIO(output))
    assert workbook.active[address].value == expected
    assert workbook.active.data_validations.dataValidation[-1].prompt == '원문 입력 안내'
    workbook.close()


def test_none_validation_with_formula_is_not_treated_as_a_prompt_only_rule(tmp_path):
    def invalid(workbook):
        rule = DataValidation(formula1='1', prompt='불명확한 규칙')
        workbook.active.add_data_validation(rule)
        rule.add('B2')
    t = target(inspect(source(tmp_path, fixture(extra=invalid))), 'B2')
    assert not t.editable and t.unsupportedReason == 'UNRESOLVED_OPTION'
