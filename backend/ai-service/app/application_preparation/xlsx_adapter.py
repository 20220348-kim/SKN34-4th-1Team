"""XLSX cells are native addresses; only proven blank inputs may receive facts."""
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from copy import deepcopy
import posixpath
import re
import warnings
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

from app.application_preparation.document_contract import (
    DocumentError, DocumentMap, ENGINES, MAX_BYTES, NativeTarget,
    NativeTargetAnalysis, WritePlan, digest,
)

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MAX_CELLS = 100000
MAIN_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
X14 = "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}"
XM = "{http://schemas.microsoft.com/office/excel/2006/main}"


def _comment_only_vml(value: bytes) -> bool:
    try:
        root = ET.fromstring(value)
        shapes = [node for node in root.iter() if node.tag == "{urn:schemas-microsoft-com:vml}shape"]
        return bool(shapes) and all(
            shape.get("type") == "#_x0000_t202" and
            (client := shape.find("{urn:schemas-microsoft-com:office:excel}ClientData")) is not None and
            client.get("ObjectType") == "Note" for shape in shapes)
    except ET.ParseError:
        return False


def _package(data: bytes) -> dict[str, bytes]:
    if not 0 < len(data) <= MAX_BYTES:
        raise DocumentError("LIMIT_EXCEEDED")
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [e.filename for e in entries]
            if (len(entries) > 512 or len(set(names)) != len(names)
                    or sum(e.file_size for e in entries) > MAX_BYTES
                    or any(e.flag_bits & 1 or e.filename.startswith(("/", "\\"))
                           or ".." in e.filename.replace("\\", "/").split("/") for e in entries)):
                raise DocumentError("LIMIT_EXCEEDED")
            parts = {e.filename: archive.read(e) for e in entries}
        for name, value in parts.items():
            if name.endswith((".xml", ".rels")):
                normalized = value.replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in normalized or b"<!ENTITY" in normalized:
                    raise DocumentError("UNSUPPORTED", reason="XLSX_UNSAFE_XML")
                ET.fromstring(value)
        types = ET.fromstring(parts["[Content_Types].xml"])
        if not any(n.get("PartName") == "/xl/workbook.xml" and n.get("ContentType") == MAIN_TYPE for n in types):
            raise DocumentError("UNSUPPORTED", reason="XLSX_ONLY")
        if any(any(marker in content_type for marker in
                       ("macroenabled", "vbaproject", "activex", "oleobject", "chart", "controlproperties", "digital-signature"))
                   or "drawing" in content_type and content_type != "application/vnd.openxmlformats-officedocument.vmldrawing"
               for n in types if (content_type := n.get("ContentType", "").lower())):
            raise DocumentError("UNSUPPORTED", reason="XLSX_UNSUPPORTED_OBJECT_OR_EXTERNAL_LINK")
        # Never sign, bypass, evaluate, or silently remove unsupported workbook objects.
        if any(name.startswith(("xl/externalLinks/", "_xmlsignatures/", "xl/activeX/",
                                "xl/ctrlProps/", "xl/embeddings/", "xl/charts/"))
               or name.startswith("xl/drawings/") and not (name.endswith(".vml") and _comment_only_vml(parts[name]))
               or "vba" in name.lower() for name in parts):
            raise DocumentError("UNSUPPORTED", reason="XLSX_UNSUPPORTED_OBJECT_OR_EXTERNAL_LINK")
        for name, value in parts.items():
            if name.endswith(".rels"):
                root = ET.fromstring(value)
                if any(n.get("Type", "").lower().rsplit("/", 1)[-1] in {
                        "vbaproject", "oleobject", "control", "drawing", "chart",
                        "activexcontrol", "externallink", "externallinkpath", "customui"} for n in root):
                    raise DocumentError("UNSUPPORTED", reason="XLSX_UNSUPPORTED_OBJECT_OR_EXTERNAL_LINK")
                if any(n.get("TargetMode") == "External" and not n.get("Type", "").endswith("/hyperlink") for n in root):
                    raise DocumentError("UNSUPPORTED", reason="XLSX_EXTERNAL_RELATIONSHIP")
        return parts
    except DocumentError:
        raise
    except (KeyError, ValueError, OSError, RuntimeError, zipfile.BadZipFile, ET.ParseError):
        raise DocumentError("UNSUPPORTED", reason="XLSX_INVALID_PACKAGE") from None


def _sheet_parts(parts: dict[str, bytes]) -> dict[str, str]:
    relations = {n.get("Id"): n.get("Target", "") for n in ET.fromstring(parts["xl/_rels/workbook.xml.rels"])}
    result = {}
    for sheet in ET.fromstring(parts["xl/workbook.xml"]).findall("./" + S + "sheets/" + S + "sheet"):
        target = relations.get(sheet.get(R + "id"), "")
        name = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
        if not name.startswith("xl/worksheets/") or name not in parts or sheet.get("name") in result:
            raise DocumentError("UNSUPPORTED", reason="XLSX_INVALID_SHEET_RELATIONSHIP")
        result[sheet.get("name")] = name
    if not 1 <= len(result) <= 30:
        raise DocumentError("LIMIT_EXCEEDED")
    return result


def _extended_validations(parts: dict[str, bytes], paths: dict[str, str]) -> dict[str, list[DataValidation]]:
    result = {}
    for sheet_name, path in paths.items():
        root = ET.fromstring(parts[path])
        rules = []
        extensions = root.find(S + "extLst")
        for extension in extensions if extensions is not None else []:
            if len(extension) != 1 or extension[0].tag != X14 + "dataValidations":
                raise DocumentError("UNSUPPORTED", reason="XLSX_EXTENDED_VALIDATION_UNSUPPORTED")
            group = extension[0]
            items = group.findall(X14 + "dataValidation")
            if len(items) != len(group) or group.get("count") != str(len(items)):
                raise DocumentError("UNSUPPORTED", reason="XLSX_EXTENDED_VALIDATION_UNSUPPORTED")
            for item in items:
                formula = item.find("./" + X14 + "formula1/" + XM + "f")
                address = item.find(XM + "sqref")
                if item.get("type") != "list" or formula is None or not formula.text or address is None or not address.text:
                    raise DocumentError("UNSUPPORTED", reason="XLSX_EXTENDED_VALIDATION_UNSUPPORTED")
                try:
                    rule = DataValidation(type="list", formula1=formula.text)
                    rule.add(address.text)
                except (TypeError, ValueError):
                    raise DocumentError("UNSUPPORTED", reason="XLSX_EXTENDED_VALIDATION_UNSUPPORTED") from None
                rules.append(rule)
        result[sheet_name] = rules
    return result


def _workbook(data: bytes, parts: dict[str, bytes]):
    paths = _sheet_parts(parts)
    extended = _extended_validations(parts, paths)
    total = 0
    for path in paths.values():
        root = ET.fromstring(parts[path])
        cells = root.findall("./" + S + "sheetData/" + S + "row/" + S + "c")
        if len({n.get("r") for n in cells}) != len(cells):
            raise DocumentError("UNSUPPORTED", reason="XLSX_DUPLICATE_CELL")
        if any(not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", n.get("r", "")) for n in cells):
            raise DocumentError("UNSUPPORTED", reason="XLSX_INVALID_CELL_ADDRESS")
        bounds = [range_boundaries(n.get("r")) for n in cells]
        rows = max((b[3] for b in bounds), default=1)
        columns = max((b[2] for b in bounds), default=1)
        for node in root.findall("./" + S + "mergeCells/" + S + "mergeCell"):
            b = range_boundaries(node.get("ref"))
            rows, columns = max(rows, b[3]), max(columns, b[2])
        total += rows * columns
        if total > MAX_CELLS or rows > 10000 or columns > 256:
            raise DocumentError("LIMIT_EXCEEDED", reason="XLSX_CELL_BUDGET")
        dimension = root.find(S + "dimension")
        if dimension is not None:
            b = range_boundaries(dimension.get("ref"))
            if b[3] * b[2] > MAX_CELLS:
                raise DocumentError("LIMIT_EXCEEDED", reason="XLSX_USED_RANGE_BUDGET")
    try:
        with warnings.catch_warnings(record=True) as notices:
            warnings.simplefilter("always")
            workbook = load_workbook(BytesIO(data), data_only=False, keep_links=False, rich_text=True)
        if any(str(notice.message) != "Data Validation extension is not supported and will be removed"
               or not any(extended.values()) for notice in notices):
            workbook.close()
            raise DocumentError("UNSUPPORTED", reason="XLSX_READER_WARNING")
        return workbook, paths, extended
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("UNSUPPORTED", reason="XLSX_REOPEN_FAILED") from None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _options(workbook, sheet, validation) -> list[str] | None:
    if validation.type != "list" or not validation.formula1:
        return None
    source = validation.formula1
    if source.startswith('"') and source.endswith('"'):
        values = source[1:-1].split(",")
    else:
        source = source.removeprefix("=")
        destinations = []
        named = sheet.defined_names.get(source) or workbook.defined_names.get(source)
        if named is not None:
            try:
                destinations = list(named.destinations)
            except (ValueError, AttributeError):
                return None
        else:
            match = re.fullmatch(r"(?:(?:'((?:[^']|'')+)'|([^'!]+))!)?(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", source)
            if match:
                destinations = [((match[1] or match[2] or sheet.title).replace("''", "'"), match[3])]
        if len(destinations) != 1:
            return None
        name, address = destinations[0]
        if name not in workbook.sheetnames:
            return None
        bounds = range_boundaries(address)
        if (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1) > 200:
            return None
        cells = [c for row in workbook[name].iter_rows(min_row=bounds[1], max_row=bounds[3],
                 min_col=bounds[0], max_col=bounds[2]) for c in row]
        if any(c.data_type == "f" or c.value is not None and not isinstance(c.value, str) for c in cells):
            return None
        values = [_text(c.value) for c in cells if c.value is not None]
    if not values or any(not v or len(v) > 1000 for v in values):
        return None
    return list(dict.fromkeys(values))


def cell_value(target: NativeTarget, value: str):
    """Keep identifiers textual; convert only explicit numeric/date formats."""
    locator = target.nativeLocator
    validation = locator.get("dataValidation")
    # A type=none rule without formulas carries an input prompt, not an option constraint.
    constrained = bool(validation and not (
        validation.get("promptOnly") is True and validation.get("type") is None
        and not validation.get("formula1") and not validation.get("formula2")))
    if constrained and (not validation.get("allowedValues") or value not in validation["allowedValues"]):
        raise DocumentError("MAPPING_FAILED", reason="UNRESOLVED_OPTION")
    if len(value) > 6000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise DocumentError("UNSUPPORTED", reason="XLSX_INVALID_VALUE")
    label = " ".join(locator.get("fieldLabels", []))
    number_format = locator["numberFormat"]
    if constrained or number_format == "@" or re.search(r"사업자|등록번호|전화|연락처|우편|이메일|코드|번호", label):
        return value
    if is_date_format(number_format):
        try:
            parsed = datetime.fromisoformat(value)
            if parsed.tzinfo is not None:
                raise ValueError()
            return parsed
        except ValueError:
            raise DocumentError("UNSUPPORTED", reason="XLSX_DATE_VALUE_REQUIRED") from None
    # General cells default to text: no guessing from the model's answer.
    if number_format == "General" or not re.search(r"[0#?]", number_format):
        return value
    try:
        percent = "%" in number_format
        number = Decimal(value.removesuffix("%"))
        if not number.is_finite() or len(number.as_tuple().digits) > 15 or abs(number) > Decimal("1e100"):
            raise InvalidOperation()
        if percent:
            if not value.endswith("%"):
                raise InvalidOperation()
            number /= 100
        elif value.endswith("%"):
            raise InvalidOperation()
        return float(number) if number != number.to_integral_value() else int(number)
    except (ValueError, InvalidOperation, OverflowError):
        raise DocumentError("UNSUPPORTED", reason="XLSX_NUMERIC_VALUE_REQUIRED") from None


class XlsxDocumentAdapter:
    def _inspect_bytes(self, data: bytes):
        parts = _package(data)
        workbook, paths, extended = _workbook(data, parts)
        sparse_layout = sum(len(ET.fromstring(parts[path]).findall("./" + S + "sheetData/" + S + "row/" + S + "c"))
                            for path in paths.values()) > 3000
        targets = []
        sheets = []
        workbook_protected = bool(workbook.security and
                                  (workbook.security.lockStructure or workbook.security.lockWindows))
        for sheet in workbook:
            root = ET.fromstring(parts[paths[sheet.title]])
            if not re.search(rb'<worksheet\b[^>]*xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"', parts[paths[sheet.title]]):
                workbook.close()
                raise DocumentError("UNSUPPORTED", reason="XLSX_NAMESPACE_SERIALIZATION")
            nodes = {n.get("r"): n for n in root.findall("./" + S + "sheetData/" + S + "row/" + S + "c")}
            merged = {}
            for region in sheet.merged_cells.ranges:
                for row in range(region.min_row, region.max_row + 1):
                    for col in range(region.min_col, region.max_col + 1):
                        merged[(row, col)] = region
            def master(row, col):
                if row < 1 or col < 1:
                    return None
                region = merged.get((row, col))
                return sheet.cell(region.min_row, region.min_col) if region else sheet.cell(row, col)
            def concealed(cell):
                return (sheet.sheet_state != "visible" or bool(sheet.row_dimensions[cell.row].hidden)
                        or any(d.hidden and d.min <= cell.column <= d.max for d in sheet.column_dimensions.values()))
            def label(cell):
                if cell is not None and concealed(cell):
                    return ""
                if cell is None or cell.data_type != "s" or not isinstance(cell.value, (str, CellRichText)):
                    return ""
                value = _text(cell.value).strip()
                return value if (0 < len(value) <= 150 and re.search(r"[a-zA-Z가-힣]", value)
                                 and not re.search(r"[□☐☑☒]", value)
                                 and not value.startswith(("※", "*", "ex)", "예)"))) else ""
            formulas = sum(c.data_type == "f" for row in sheet for c in row)
            sheet_info = {"sheetName": sheet.title, "state": sheet.sheet_state,
                          "usedRange": sheet.calculate_dimension(), "mergedRanges": sorted(map(str, sheet.merged_cells.ranges)),
                          "freezePanes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                          "tables": [{"name": t.name, "ref": t.ref} for t in sheet.tables.values()],
                          "dataValidations": [{"type": v.type, "sourceRange": str(v.sqref),
                              "formula1": v.formula1, "formula2": v.formula2} for v in
                              [*sheet.data_validations.dataValidation, *extended[sheet.title]]],
                          "formulaCount": formulas, "protected": bool(sheet.protection.sheet)}
            sheets.append(sheet_info)
            section = ""
            for row in sheet:
                for cell in row:
                    region = merged.get((cell.row, cell.column))
                    child = isinstance(cell, MergedCell)
                    if not child and region and region.min_row == cell.row and cell.value and region.max_col - region.min_col >= 3:
                        section = label(cell) or section
                    left_cell = master(cell.row, cell.column - 1)
                    above_cell = master(cell.row - 1, cell.column)
                    left = label(left_cell) if left_cell is not cell else ""
                    above = label(above_cell) if above_cell is not cell else ""
                    # A merged title immediately above a row is section context, not a column input label.
                    above_region = merged.get((cell.row - 1, cell.column))
                    if above_region and above_region.max_col - above_region.min_col >= 3:
                        above = ""
                    left_region = merged.get((cell.row, cell.column - 1))
                    if left_region and left_region.max_col >= cell.column:
                        left = ""
                    # Styled header rows disambiguate repeated people/product rows from
                    # example values above. Only the first blank slot below a header is eligible.
                    header = ""
                    header_row = None
                    for previous_row in range(cell.row - 1, max(0, cell.row - 9), -1):
                        head = master(previous_row, cell.column)
                        if head and head.font.bold and label(head) and sum(
                                bool(label(c)) for c in sheet[previous_row]) >= 3:
                            header, header_row = label(head), previous_row
                            break
                    if header:
                        above = header
                        left = ""  # a neighboring person's name is not this column's field label
                    labels = list(dict.fromkeys(v for v in (left, above) if v))
                    repeated_blank = bool(header_row and cell.row > header_row + 1
                                          and sheet.cell(cell.row - 1, cell.column).value is None)
                    validations = [v for v in [*sheet.data_validations.dataValidation, *extended[sheet.title]]
                                   if cell.coordinate in v]
                    validation = None
                    if validations:
                        v = validations[0]
                        allowed = _options(workbook, sheet, v) if len(validations) == 1 else None
                        validation = {"type": v.type, "formula1": v.formula1, "formula2": v.formula2,
                                      "sourceRange": str(v.sqref), "allowedValues": allowed,
                                      "displayValues": allowed, "allowBlank": v.allowBlank,
                                      "promptOnly": len(validations) == 1 and v.type is None
                                      and not v.formula1 and not v.formula2}
                    hidden = concealed(cell)
                    protected = bool(workbook_protected or sheet.protection.sheet and cell.protection.locked)
                    node = nodes.get(cell.coordinate)
                    complex_cell = node is not None and (any(n.tag not in {S + "v", S + "is"} for n in node)
                                   or any(k not in {"r", "s", "t"} for k in node.attrib))
                    table_region = None
                    for table in sheet.tables.values():
                        b = range_boundaries(table.ref)
                        if b[0] <= cell.column <= b[2] and b[1] <= cell.row <= b[3]:
                            table_region = table.ref
                            break
                    blank = cell.value is None
                    evidence = bool(labels and node is not None)
                    reason = ("FORMULA_CELL" if cell.data_type == "f" else "MERGED_CHILD" if child else
                              "HIDDEN_CELL" if hidden else "PROTECTED_CELL" if protected else
                              "DATA_TABLE" if table_region else "REPEATED_ROW_NOT_SELECTED" if repeated_blank else "UNSUPPORTED_CELL_OBJECT" if complex_cell else
                               "UNRESOLVED_OPTION" if validation and not validation["promptOnly"] and not validation["allowedValues"] else
                              "NONEMPTY_CELL" if not blank else "AMBIGUOUS_BLANK_CELL" if not evidence else None)
                    # Styled blank cells without a label or native validation are layout, not inputs.
                    if blank and not child and (node is None or sparse_layout and not labels and not validations):
                        continue
                    text = _text(cell.value)
                    if len(text) > 6000 or len(targets) >= 3000:
                        workbook.close()
                        raise DocumentError("LIMIT_EXCEEDED")
                    editable = reason is None
                    classification = "DATA_TABLE" if table_region else "FORM_TABLE" if evidence else "AMBIGUOUS"
                    locator = {"sheetName": sheet.title, "cellAddress": cell.coordinate, "row": cell.row,
                               "column": cell.column, "value": cell.value if isinstance(cell.value, (str, int, float, bool)) else
                               text if cell.value is not None else None,
                               "displayValue": text, "displayPolicy": "RAW_WITH_NUMBER_FORMAT",
                               "dataType": cell.data_type, "cellType": "merged_child" if child else
                               "formula" if cell.data_type == "f" else "blank" if blank else
                               "date" if cell.data_type == "d" else "boolean" if cell.data_type == "b" else
                               "numeric" if cell.data_type == "n" else "constant_text", "numberFormat": cell.number_format,
                               "formula": cell.data_type == "f", "merged": region is not None,
                               "mergedMaster": sheet.cell(region.min_row, region.min_col).coordinate if region else None,
                               "mergedRange": str(region) if region else None, "hidden": hidden,
                               "protected": protected, "locked": bool(cell.protection.locked), "editable": editable,
                               "bindingEligible": editable, "fieldLabels": labels, "rowLabels": [left] if left else [],
                               "columnLabels": [above] if above else [], "tableHeadings": [section] if section else [],
                               "sectionPath": [sheet.title, section] if section else [sheet.title],
                               "regionClassification": "DATA_TABLE" if table_region else "FORM_REGION" if evidence else
                               "SUMMARY_REGION" if cell.data_type == "f" else "AMBIGUOUS",
                               "tableRegion": table_region, "dataValidation": validation}
                    targets.append(NativeTarget(targetId=f"xlsx:s:{quote(sheet.title, safe='')}:c:{cell.coordinate}",
                        kind="XLSX_CELL", currentText=text, label=" / ".join(labels), nativeLocator=locator,
                        context=(" | ".join([sheet.title, section, left, above]))[:1000], editable=editable,
                        unsupportedReason=reason, analysis=NativeTargetAnalysis(semanticSection=section or sheet.title,
                            sectionPath=locator["sectionPath"], tableClassification=classification,
                            tableClassificationConfidence=0.8 if evidence else 0.3, reviewRequired=not evidence)))
        document = DocumentMap(sourceSha256=digest(data), format="xlsx", engineVersion=ENGINES["xlsx"],
                               targets=targets, workbookMetadata={"sheets": sheets,
                               "protected": workbook_protected})
        return document, workbook, parts, paths

    async def inspect(self, path: Path) -> DocumentMap:
        document, workbook, _, _ = self._inspect_bytes(path.read_bytes())
        workbook.close()
        return document

    async def apply(self, path: Path, document: DocumentMap, plan: WritePlan, facts: dict[str, str]):
        source = path.read_bytes()
        if (digest(source) != document.sourceSha256 or plan.sourceSha256 != document.sourceSha256
                or plan.mapVersion != document.mapVersion or document.engineVersion != ENGINES["xlsx"]):
            raise DocumentError("SOURCE_CHANGED")
        fresh, workbook, parts, paths = self._inspect_bytes(source)
        before = {t.targetId: t for t in fresh.targets}
        edits = {}
        expected = {}
        try:
            for op in plan.operations:
                target = before.get(op.targetId)
                if (target is None or not target.editable or target.currentText != op.expectedText
                        or op.targetId not in plan.scopeTargetIds):
                    raise DocumentError("MAPPING_FAILED", reason="TARGET_NOT_EDITABLE_OR_OUT_OF_SCOPE")
                if (op.operation not in {"input", "set_field"} or op.start != 0 or op.end != 0 or op.box is not None
                        or op.targetId in edits or op.valueRef not in facts):
                    raise DocumentError("UNSUPPORTED", reason="XLSX_ONLY_BLANK_CELL_WRITE")
                value = cell_value(target, facts[op.valueRef])
                sheet_name, address = target.nativeLocator["sheetName"], target.nativeLocator["cellAddress"]
                cell = workbook[sheet_name][address]
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"  # "=..." must remain literal user text, never a new formula.
                edits[op.targetId] = (sheet_name, address)
                expected[op.targetId] = value
            staging = BytesIO()
            workbook.save(staging)
        finally:
            workbook.close()
        # openpyxl serializes values. Graft only those c elements into original worksheet
        # XML, retaining all other parts and worksheet structure (including cached formulas).
        with zipfile.ZipFile(staging) as saved:
            saved_parts = {name: saved.read(name) for name in saved.namelist()}
        saved_paths = _sheet_parts(saved_parts)
        changed = {}
        for target_id, (sheet_name, address) in edits.items():
            part = paths[sheet_name]
            original_xml = changed.get(part, parts[part])
            # Native XLSX transitional default namespace only; no guessed prefixed addresses.
            root = ET.fromstring(original_xml)
            if not re.search(rb'<worksheet\b[^>]*xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"', original_xml):
                raise DocumentError("UNSUPPORTED", reason="XLSX_NAMESPACE_SERIALIZATION")
            saved_root = ET.fromstring(saved_parts[saved_paths[sheet_name]])
            new = saved_root.find(f"./{S}sheetData/{S}row/{S}c[@r='{address}']")
            old = root.find(f"./{S}sheetData/{S}row/{S}c[@r='{address}']")
            if new is None or old is None:
                raise DocumentError("VALIDATION_FAILED", reason="XLSX_CELL_NOT_MATERIALIZED")
            graft = deepcopy(new)
            graft.attrib = {**old.attrib, "t": new.get("t", "n")}
            # Strip serializer namespace prefixes only on the new, value-only cell.
            serialized = ET.tostring(graft, encoding="utf-8")
            serialized = re.sub(rb'<(/?)ns\d+:', rb'<\1', serialized)
            serialized = re.sub(rb' xmlns:ns\d+="[^"]+"', b"", serialized)
            pattern = rb'<c\b(?=[^>]*\br="' + address.encode() + rb'")[^>]*?(?:/>|>.*?</c>)'
            matches = list(re.finditer(pattern, original_xml, re.DOTALL))
            if len(matches) != 1:
                raise DocumentError("UNSUPPORTED", reason="XLSX_CELL_XML_AMBIGUOUS")
            match = matches[0]
            changed[part] = original_xml[:match.start()] + serialized + original_xml[match.end():]
        buffer = BytesIO()
        with zipfile.ZipFile(BytesIO(source)) as original, zipfile.ZipFile(buffer, "w") as output:
            for entry in original.infolist():
                output.writestr(entry, changed.get(entry.filename, parts[entry.filename]))
        result = buffer.getvalue()
        after_document, reopened, after_parts, _ = self._inspect_bytes(result)
        try:
            if fresh.workbookMetadata != after_document.workbookMetadata:
                raise DocumentError("VALIDATION_FAILED", reason="XLSX_WORKBOOK_STRUCTURE_CHANGED")
            if parts.keys() != after_parts.keys() or any(after_parts[n] != v for n, v in parts.items() if n not in changed):
                raise DocumentError("VALIDATION_FAILED", reason="XLSX_UNEDITED_PART_CHANGED")
            # Remove approved cell payloads and compare complete XML structures, including
            # every formula, other value, row/column dimension and validation/CF node.
            for part in changed:
                old_root, new_root = ET.fromstring(parts[part]), ET.fromstring(after_parts[part])
                for sheet_name, address in edits.values():
                    if paths[sheet_name] != part:
                        continue
                    for root in (old_root, new_root):
                        node = root.find(f"./{S}sheetData/{S}row/{S}c[@r='{address}']")
                        node.attrib.pop("t", None)
                        for child in list(node):
                            node.remove(child)
                if ET.tostring(old_root) != ET.tostring(new_root):
                    raise DocumentError("VALIDATION_FAILED", reason="XLSX_STRUCTURE_OR_STYLE_CHANGED")
            for target_id, (sheet_name, address) in edits.items():
                if reopened[sheet_name][address].value != expected[target_id]:
                    raise DocumentError("VALIDATION_FAILED", reason="XLSX_TARGET_VALUE_MISMATCH")
        finally:
            reopened.close()
        if digest(path.read_bytes()) != document.sourceSha256:
            raise DocumentError("SOURCE_CHANGED")
        return result, {"requested": len(edits), "resolved": len(edits), "applied": len(edits),
                        "verified": len(edits), "unresolved": 0, "reopened": True, "xml": "PASSED",
                        "styleStructure": "PASSED", "formulas": "PASSED", "dataValidation": "PASSED",
                        "unchangedParts": "PASSED", "render": "NOT_RUN"}
