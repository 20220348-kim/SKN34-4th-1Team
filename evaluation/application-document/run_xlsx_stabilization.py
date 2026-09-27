"""One additional official XLSX: safe-input mapping, unsafe-scope audit and native write."""
import argparse
import asyncio
import base64
from collections import Counter
import json
import os
from pathlib import Path
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend/ai-service"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.application_preparation.document_contract import (
    DocumentFieldReference, DocumentError, EditOperation, GenerateDocumentRequest,
    MapDocumentRequest, PlanSelection, digest, mapping_label_matches, validate_mapping, validate_plan,
)
from app.application_preparation.document_pipeline import inspect_document
from app.application_preparation.xlsx_adapter import XlsxDocumentAdapter, cell_value
from evaluate_agent_mapping import ContextAgent
from run_docx_mapping import RecordedModel


def snapshot(path):
    workbook = load_workbook(path, data_only=False, rich_text=True)
    sheets = []
    for sheet in workbook:
        formulas = [{"address": c.coordinate, "expression": c.value, "styleId": c.style_id}
                    for row in sheet for c in row if c.data_type == "f"]
        sheets.append({"name": sheet.title, "state": sheet.sheet_state,
            "usedRange": sheet.calculate_dimension(), "mergedRanges": sorted(map(str, sheet.merged_cells.ranges)),
            "formulas": formulas, "freezePanes": sheet.freeze_panes,
            "rowHeights": {str(k): v.height for k, v in sheet.row_dimensions.items() if v.height is not None},
            "columnWidths": {k: v.width for k, v in sheet.column_dimensions.items()},
            "hiddenRows": [k for k, v in sheet.row_dimensions.items() if v.hidden],
            "hiddenColumns": [k for k, v in sheet.column_dimensions.items() if v.hidden],
            "protection": ET.tostring(sheet.protection.to_tree(), encoding="unicode"),
            "lockedCellCount": sum(c.protection.locked for row in sheet for c in row),
            "unlockedCellCount": sum(not c.protection.locked for row in sheet for c in row),
            "validations": [ET.tostring(v.to_tree(), encoding="unicode") for v in sheet.data_validations.dataValidation],
            "conditionalFormattingCount": len(sheet.conditional_formatting)})
    result = {"sheetCount": len(workbook.sheetnames), "sheetNames": workbook.sheetnames,
        "workbookProtection": ET.tostring(workbook.security.to_tree(), encoding="unicode") if workbook.security else None,
        "definedNames": {k: ET.tostring(v.to_tree(), encoding="unicode") for k, v in workbook.defined_names.items()},
        "sheets": sheets}
    workbook.close()
    return result


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


async def run(args):
    data = args.source.read_bytes()
    truth = json.loads(args.expected.read_text(encoding="utf-8"))
    blind = json.loads(args.blind.read_text(encoding="utf-8"))
    if digest(data) != truth["sourceSha256"] or blind["sourceSha256"] != truth["sourceSha256"]:
        raise ValueError("Official source SHA mismatch")
    if not 5 <= len(truth["fields"]) <= 10 or "expectedTargetId" in json.dumps(blind):
        raise ValueError("Ground truth size or blind-input boundary invalid")
    request = MapDocumentRequest(sourceBase64=base64.b64encode(data).decode(), sourceSha256=digest(data),
        format="xlsx", scope=blind["scopeTitle"], fields=[DocumentFieldReference(**f) for f in blind["fields"]])
    document = await inspect_document(args.source, request)
    targets = {t.targetId: t for t in document.targets}
    candidates = {f.id: [t.targetId for t in document.targets if t.editable
        and t.nativeLocator.get("bindingEligible") and mapping_label_matches(f.label, t, f.guidance)]
        for f in request.fields}
    writable = [f for f in truth["fields"] if f["expectedWritable"]]
    manual = [f for f in truth["fields"] if not f["expectedWritable"]]
    if any(f["expectedTargetId"] not in candidates[f["id"]] for f in writable):
        raise ValueError("Safe input candidate coverage incomplete")
    if any(targets[f["expectedTargetId"]].editable for f in manual):
        raise ValueError("Formula safety probe unexpectedly editable")
    unsafe = [t for t in document.targets if t.nativeLocator.get("formula") or t.nativeLocator.get("hidden")
              or t.nativeLocator.get("protected") or t.unsupportedReason == "MERGED_CHILD"]
    if any(t.editable for t in unsafe):
        raise ValueError("Unsafe source cells exposed as candidates")
    args.output.mkdir(parents=True, exist_ok=True)
    save(args.output / "preflight.json", {"sourceSha256": digest(data), "engineVersion": document.engineVersion,
        "safeInputFields": len(writable), "unsafeOutputProbes": len(manual),
        "safeInputCoverage": 1.0, "candidates": candidates, "unsafeEditableCount": 0,
        "unsafeReasons": dict(Counter(t.unsupportedReason for t in unsafe)),
        "structure": snapshot(args.source), "apiCalls": 0})
    if args.mapping:
        attempt = args.output / "mapping-attempt.json"
        previous = list((Path(__file__).parent / "runs").glob("xlsx-*/mapping-attempt.json"))
        if attempt.exists() or any(json.loads(p.read_text(encoding="utf-8")).get("sourceSha256") == digest(data)
                                   for p in previous):
            raise FileExistsError("Mapping already attempted for this source")
        if len(previous) != 2:
            raise ValueError("This stabilization allows exactly one new source after the two initial calls")
        from dotenv import load_dotenv
        from langchain_openai import ChatOpenAI
        load_dotenv(Path(os.environ["XLSX_EVALUATION_ENV"]), override=False)
        if not os.getenv("OPENAI_API_KEY"):
            raise ValueError("OpenAI key unavailable")
        agent = ContextAgent(evidence=None, model=RecordedModel(ChatOpenAI(model=args.model,
            api_key=os.environ["OPENAI_API_KEY"], use_responses_api=True, store=False,
            reasoning={"effort": "none"}, max_retries=0), args.output), run_timeout_seconds=240)
        record = {"sourceSha256": digest(data), "status": "MODEL_CALL_STARTED", "maximumApiCalls": 1,
                  "authorization": "One additional official XLSX in stabilization r2", "model": args.model}
        save(attempt, record)
        try:
            selection = await agent.map_document(request, document)
            validate_mapping(request, document, selection)
            by_field = {}
            for binding in selection.bindings:
                by_field.setdefault(binding.factId, []).append(binding.targetId)
            correct = sum(by_field.get(f["id"]) == [f["expectedTargetId"]] for f in writable)
            wrong = sum(bool(by_field.get(f["id"])) and by_field[f["id"]] != [f["expectedTargetId"]] for f in writable)
            unsafe_bindings = sum(bool(by_field.get(f["id"])) for f in manual)
            save(args.output / "mapping.json", {"sourceSha256": digest(data), "engineVersion": document.engineVersion,
                "apiCalls": agent.calls, "model": args.model, "productionValidation": "PASS",
                "correctSafeInputs": correct, "wrongSafeInputs": wrong,
                "unmappedSafeInputs": len(writable) - correct - wrong, "expectedReadonlyUnmapped": len(manual),
                "actualReadonlyUnmapped": sum(f["id"] in selection.unmappedFieldIds for f in manual),
                "unsafeBindings": unsafe_bindings, "wrongTargetRate": (wrong + unsafe_bindings) / len(truth["fields"]),
                "bindings": [b.model_dump() for b in selection.bindings], "scopeTargetIds": selection.scopeTargetIds,
                "unmappedFieldIds": selection.unmappedFieldIds,
                "groundTruthUsage": "Expected IDs/writability withheld; calculated outputs are explicit safety probes",
                "humanReview": "AGENT_REVIEWED_NOT_INDEPENDENT_USER_REVIEWED"})
            record.update(status="COMPLETED", apiCalls=agent.calls)
        except Exception as error:
            record.update(status="MODEL_CALL_FAILED", apiCalls=agent.calls, errorType=type(error).__name__)
            raise
        finally:
            save(attempt, record)
    if args.write:
        if args.write.resolve() == args.source.resolve():
            raise ValueError("Source must not be overwritten")
        facts = {f["id"]: f["writeValue"] for f in writable}
        generation = GenerateDocumentRequest(sourceBase64=base64.b64encode(data).decode(), sourceSha256=digest(data),
            format="xlsx", answerRevision=1, scope=truth["scopeTitle"],
            facts=[{"id": f["id"], "label": f["label"], "value": f["writeValue"]} for f in writable])
        operations = [EditOperation(targetId=f["expectedTargetId"], operation="input", expectedText="",
            start=0, end=0, valueRef=f["id"], box=None, reason="Reviewed unlocked official input, dummy value") for f in writable]
        plan = validate_plan(generation, document, PlanSelection(operations=operations, unresolvedTargets=[],
            scopeTargetIds=[op.targetId for op in operations]))
        output, verification = await XlsxDocumentAdapter().apply(args.source, document, plan, facts)
        args.write.write_bytes(output)
        reopened = load_workbook(args.write, data_only=False)
        for f in writable:
            target = targets[f["expectedTargetId"]]
            assert reopened[target.nativeLocator["sheetName"]][target.nativeLocator["cellAddress"]].value == cell_value(target, f["writeValue"])
        reopened.close()
        if snapshot(args.source) != snapshot(args.write) or digest(args.source.read_bytes()) != digest(data):
            raise ValueError("Workbook structure/source changed")
        save(args.output / "write-verification.json", {"sourceSha256": digest(data), "outputSha256": digest(output),
            "writes": len(writable), "verification": verification, "structureComparison": "PASS",
            "formulaCount": sum(len(s["formulas"]) for s in snapshot(args.source)["sheets"]), "unsafeWrites": 0})
    return {"safeInputs": len(writable), "readonlyProbes": len(manual), "engineVersion": document.engineVersion}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    for name in ("source", "expected", "blind", "output"):
        parser.add_argument("--" + name, type=Path, required=True)
    parser.add_argument("--mapping", action="store_true")
    parser.add_argument("--write", type=Path)
    parser.add_argument("--model", default="gpt-5.6-luna")
    print(json.dumps(asyncio.run(run(parser.parse_args())), ensure_ascii=False, indent=2))
